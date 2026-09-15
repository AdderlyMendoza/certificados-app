"""Endpoints de generación masiva: preview de dataset, crear, lanzar, estado y ZIP."""
from __future__ import annotations

import io
from typing import Annotated

from fastapi import APIRouter, Depends, File, Form, Query, UploadFile, status
from fastapi.responses import FileResponse, StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.core.config import settings
from app.core.dependencies import CurrentUser, DbSession, require_operator
from app.core.exceptions import ValidationError
from app.schemas.common import Message, Page
from app.schemas.generation import DatasetPreview, GenerationCreate, GenerationRead
from app.services.generation_service import GenerationService

router = APIRouter(prefix="/generations", tags=["generations"])


@router.get("/template-download")
async def download_template(_: CurrentUser) -> StreamingResponse:
    """Descarga un Excel de ejemplo con las columnas requeridas."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Datos"

    headers = ["nombre_archivo", "nombres", "apellido_paterno", "apellido_materno", "cargo", "fecha", "codigo"]
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    examples = [
        ["certificado_juan_perez", "Juan", "Pérez", "García", "Ingeniero", "2026-05-15", "CERT-001"],
        ["certificado_maria_lopez", "María", "López", "Martínez", "Doctora", "2026-05-16", "CERT-002"],
        ["certificado_carlos_rivera", "Carlos", "Rivera", "Sánchez", "Licenciado", "2026-05-17", "CERT-003"],
    ]
    for row_idx, row_data in enumerate(examples, 2):
        for col_idx, value in enumerate(row_data, 1):
            ws.cell(row=row_idx, column=col_idx, value=value)

    col_widths = {"A": 30, "B": 15, "C": 18, "D": 18, "E": 18, "F": 14, "G": 14}
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    ws2 = wb.create_sheet("Instrucciones")
    instructions = [
        ["Columna", "Requerida", "Descripción"],
        ["nombre_archivo", "No", "Nombre del archivo PDF generado (sin extensión). Si está vacío, se usa 'documento_N'."],
        ["nombres", "Depende", "Nombre(s) del titular. Debe coincidir con una variable {{nombres}} en la plantilla."],
        ["apellido_paterno", "Depende", "Apellido paterno. Debe coincidir con una variable en la plantilla."],
        ["apellido_materno", "Depende", "Apellido materno. Debe coincidir con una variable en la plantilla."],
        ["cargo", "Depende", "Cargo o título. Debe coincidir con una variable en la plantilla."],
        ["fecha", "Depende", "Fecha del certificado. Debe coincidir con una variable en la plantilla."],
        ["codigo", "Depende", "Código único. Puede usarse en el QR: {{codigo}}."],
    ]
    for row_idx, row_data in enumerate(instructions, 1):
        for col_idx, value in enumerate(row_data, 1):
            cell = ws2.cell(row=row_idx, column=col_idx, value=value)
            if row_idx == 1:
                cell.fill = header_fill
                cell.font = header_font
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    ws2.column_dimensions["A"].width = 20
    ws2.column_dimensions["B"].width = 14
    ws2.column_dimensions["C"].width = 70

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=plantilla_generacion.xlsx"},
    )


@router.post(
    "/preview-dataset",
    response_model=DatasetPreview,
    dependencies=[Depends(require_operator)],
)
async def preview_dataset(
    _: CurrentUser, file: Annotated[UploadFile, File()]
) -> DatasetPreview:
    """Sube un Excel/CSV/ODS y devuelve columnas + primeras filas + errores."""
    content = await file.read()
    if len(content) > settings.max_upload_bytes:
        raise ValidationError(f"El archivo supera {settings.MAX_UPLOAD_MB} MB")
    return GenerationService.preview_dataset(content, file.filename or "dataset.xlsx")


@router.get("", response_model=Page[GenerationRead])
async def list_generations(
    db: DbSession,
    _: CurrentUser,
    page: Annotated[int, Query(ge=1)] = 1,
    size: Annotated[int, Query(ge=1, le=100)] = 20,
) -> Page[GenerationRead]:
    items, total = await GenerationService(db).list(offset=(page - 1) * size, limit=size)
    return Page(items=items, total=total, page=page, size=size)


@router.post(
    "",
    response_model=GenerationRead,
    status_code=status.HTTP_201_CREATED,
    dependencies=[Depends(require_operator)],
)
async def create_generation(
    db: DbSession,
    current_user: CurrentUser,
    name: Annotated[str, Form()],
    template_id: Annotated[int, Form()],
    file: Annotated[UploadFile, File()],
) -> GenerationRead:
    content = await file.read()
    data = GenerationCreate(name=name, template_id=template_id)
    return await GenerationService(db).create(  # type: ignore[return-value]
        data, content=content, filename=file.filename or "dataset.xlsx", owner=current_user
    )


@router.get("/{generation_id}", response_model=GenerationRead)
async def get_generation(
    db: DbSession, _: CurrentUser, generation_id: int
) -> GenerationRead:
    return await GenerationService(db).get(generation_id)  # type: ignore[return-value]


@router.post(
    "/{generation_id}/start",
    response_model=GenerationRead,
    dependencies=[Depends(require_operator)],
)
async def start_generation(
    db: DbSession, _: CurrentUser, generation_id: int
) -> GenerationRead:
    return await GenerationService(db).start(generation_id)  # type: ignore[return-value]


@router.post(
    "/{generation_id}/cancel",
    response_model=GenerationRead,
    dependencies=[Depends(require_operator)],
)
async def cancel_generation(
    db: DbSession, _: CurrentUser, generation_id: int
) -> GenerationRead:
    return await GenerationService(db).cancel(generation_id)  # type: ignore[return-value]


@router.get("/{generation_id}/download")
async def download_zip(
    db: DbSession, _: CurrentUser, generation_id: int
) -> FileResponse:
    service = GenerationService(db)
    gen = await service.get(generation_id)
    zip_path = service.zip_file(gen)
    return FileResponse(
        zip_path, filename=f"{gen.name}.zip", media_type="application/zip"
    )
